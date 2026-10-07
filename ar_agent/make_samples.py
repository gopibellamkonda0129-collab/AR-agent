"""Generate sample invoice files in several formats for manual testing."""
import datetime as dt
import os

import openpyxl
from docx import Document

SAMPLES = "sample_data"

ROWS = [
    ["customer_name", "invoice_number", "invoice_date", "due_date", "amount", "amount_paid", "status", "contact_email"],
    ["Omega Logistics", "INV-2201", "2026-05-20", "2026-06-19", 18500, 0, "open", "ap@omega-logistics.com"],
    ["Omega Logistics", "INV-2180", "2026-04-22", "2026-05-22", 9200, 0, "overdue", "ap@omega-logistics.com"],
    ["Riverside Inn", "INV-2215", "2026-07-02", "2026-08-01", 3400, 2000, "partial", "finance@riverside-inn.us"],
    ["Blue Harbor Co", "INV-2220", "2026-07-25", "2026-08-24", 760, 0, "sent", "billing@blueharbor.co"],
]


def main():
    # xlsx with a name variant header (same column ORDER as ROWS)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Client/Company", "Invoice No", "Issued", "Due", "Grand Total", "Paid", "State", "Email"])
    for r in ROWS[1:]:
        ws.append(r)
    wb.save(os.path.join(SAMPLES, "ledger_variant.xlsx"))

    # docx with a table inside the paragraph flow
    doc = Document()
    doc.add_heading("Monthly ledger", level=1)
    doc.add_paragraph("Customer invoices exported on " + dt.date.today().isoformat())
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    for i, h in enumerate(["Company", "Invoice #", "Date", "Due Date", "Balance"]):
        hdr[i].text = h
    for r in ROWS[1:]:
        cells = table.add_row().cells
        cells[0].text = r[0]
        cells[1].text = r[1]
        cells[2].text = r[2]
        cells[3].text = r[3]
        cells[4].text = str(float(r[4]) - float(r[5]))
    doc.add_paragraph("Please review the outstanding balances above.")
    doc.save(os.path.join(SAMPLES, "ledger_mix.docx"))

    # csv variant (uppercase headers, one row per invoice)
    with open(os.path.join(SAMPLES, "ledger_mix.csv"), "w", newline="", encoding="utf-8") as f:
        f.write("CUSTOMER NAME,INVOICE NUMBER,DUE DATE,AMOUNT,PAID,STATUS,EMAIL\n")
        for r in ROWS[1:]:
            f.write(f"{r[0]},{r[1]},{r[3]},{r[4]},{r[5]},{r[6]},{r[7]}\n")

    # txt delimited export
    with open(os.path.join(SAMPLES, "ledger_pipe.txt"), "w", encoding="utf-8") as f:
        f.write("Customer|Invoice|Due|Total|Paid|Status|Email\n")
        for r in ROWS[1:]:
            f.write(f"{r[0]}|{r[1]}|{r[3]}|{r[4]}|{r[5]}|{r[6]}|{r[7]}\n")


if __name__ == "__main__":
    main()
    print("Wrote sample xlsx, docx, csv, txt to", SAMPLES)